"""
routes/admin/templates_and_export.py
=====================================
Керування .docx-шаблонами документів, масова генерація документів по
групі (з фоновими job'ами прогресу), масове вивантаження фото
студентів групи.
"""
from flask import render_template, request, redirect, url_for, flash, session, send_file
from routes.db import get_db
from routes.utils import log_action, permission_required, logger
from routes.helpers import current_username, sort_ukrainian
from routes.admin import admin_bp
import sqlite3
import os
from routes.utils import get_templates_with_metadata, TEMPLATE_FOLDER
from routes.gen_docx import gen_doc
from routes import office_editor
from werkzeug.utils import secure_filename
from datetime import datetime
import zipfile
import io
import threading
import uuid
import time
from docx import Document as DocxDocument
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def _set_table_borders(table):
    """Тонка суцільна рамка й внутрішні лінії, як у зразку студквитків."""
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), 'auto')
        borders.append(el)
    tblPr.append(borders)


def _set_table_fixed_width(table, width_in, indent_in=0.0):
    """Фіксована ширина таблиці (а не автопідбір) + зсув вліво, як у зразку."""
    tblPr = table._tbl.tblPr
    tblW = OxmlElement('w:tblW')
    tblW.set(qn('w:w'), str(int(width_in * 1440)))
    tblW.set(qn('w:type'), 'dxa')
    tblPr.append(tblW)
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblPr.append(layout)
    if indent_in:
        tblInd = OxmlElement('w:tblInd')
        tblInd.set(qn('w:w'), str(int(indent_in * 1440)))
        tblInd.set(qn('w:type'), 'dxa')
        tblPr.append(tblInd)


# Ширини колонок (дюйми) - підібрані під зразок "Інформація для
# студентських квитків": №, ПІБ, Паспорт, Виданий, Дата видачі,
# Дійсний, Запис (ЄДДР), Ідентифікаційний номер.
_ID_CARD_COL_WIDTHS_IN = [0.494, 2.535, 1.108, 1.496, 0.994, 0.994, 1.369, 1.613]
_ID_CARD_HEADERS = ['№ з/п', 'ПІБ', 'Паспорт', 'Виданий', 'Дата видачі',
                     'Дійсний', 'Запис', 'Ідентифікаційний номер']


def build_id_card_info_docx(rows, group_name, degree_level, course):
    """
    Формує .docx "Інформація для студентських квитків" - така сама
    структура/форматування, як у зразку, яким Домінік користується для
    друку студентських квитків (ПВНЗ/філія зверху, таблиця з
    паспортними даними, підпис завідувача знизу).

    rows: список dict з ключами full_name, passport, issued_by,
          issue_date, valid_until, unique_number, tax_id.
    """
    doc = DocxDocument()

    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Inches(11.69), Inches(8.27)
    section.top_margin = Inches(0.25)
    section.bottom_margin = Inches(0.2)
    section.left_margin = Inches(0.79)
    section.right_margin = Inches(0.79)

    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)

    def add_centered(text, size_pt, bold):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(size_pt)
        run.bold = bold
        return p

    add_centered('ПВНЗ «Європейський університет»', 14, True)
    add_centered('Львівська філія', 14, True)
    add_centered(
        f"Інформація для студентських квитків {degree_level.lower()} {course}курс {group_name}",
        14, False
    )
    doc.add_paragraph()

    table = doc.add_table(rows=1, cols=len(_ID_CARD_HEADERS))
    table.autofit = False
    _set_table_fixed_width(table, sum(_ID_CARD_COL_WIDTHS_IN), indent_in=-0.3)
    _set_table_borders(table)

    header_cells = table.rows[0].cells
    for i, text in enumerate(_ID_CARD_HEADERS):
        header_cells[i].width = Inches(_ID_CARD_COL_WIDTHS_IN[i])
        p = header_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.bold = True

    for i, col in enumerate(table.columns):
        col.width = Inches(_ID_CARD_COL_WIDTHS_IN[i])

    for idx, r in enumerate(rows, start=1):
        cells = table.add_row().cells
        values = [
            str(idx), r['full_name'], r['passport'], r['issued_by'],
            r['issue_date'], r['valid_until'], r['unique_number'], r['tax_id'],
        ]
        for i, val in enumerate(values):
            cells[i].width = Inches(_ID_CARD_COL_WIDTHS_IN[i])
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if i != 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(val or '')
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)

    doc.add_paragraph()
    add_centered(
        'Завідувач ВУП Львівської філії' + ' ' * 35 + 'Людмила СТАНІЩУК',
        12, False
    )

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


@admin_bp.route('/admin/templates', methods=['GET', 'POST'])
@permission_required('manage_templates')
def manage_templates():
    """
    Сторінка управління Word-шаблонами (папка template_word/): перегляд
    списку, завантаження нового шаблону з необов'язковим описом і
    позначкою "тільки для адміністратора".
    """
    conn = get_db()

    if request.method == 'POST':
        file = request.files.get('template_file')
        display_name = (request.form.get('display_name') or '').strip()
        description = (request.form.get('description') or '').strip()
        admin_only = 1 if request.form.get('admin_only') == 'on' else 0
        # Галочка "Показувати в списках" (за замовчуванням увімкнена у формі):
        # знята галочка = шаблон прихований зі списків вибору при генерації.
        hidden = 0 if request.form.get('visible') == 'on' else 1

        if not file or file.filename == '':
            flash('Оберіть файл шаблону', 'danger')
            return redirect(url_for('admin.manage_templates'))

        filename = secure_filename(file.filename)
        if not filename.lower().endswith('.docx'):
            flash('Шаблон повинен бути файлом .docx', 'danger')
            return redirect(url_for('admin.manage_templates'))

        # Якщо назву не вказано - у списках вибору показується ім'я файлу
        if not display_name:
            display_name = filename

        os.makedirs(TEMPLATE_FOLDER, exist_ok=True)
        dest_path = os.path.join(TEMPLATE_FOLDER, filename)
        is_replace = os.path.exists(dest_path)

        try:
            file.save(dest_path)
            conn.execute("""
                INSERT INTO document_templates (filename, display_name, description, admin_only, hidden, uploaded_by, uploaded_at)
                VALUES (?, ?, ?, ?, ?, ?, datetime('now'))
                ON CONFLICT(filename) DO UPDATE SET
                    display_name = excluded.display_name,
                    description = excluded.description,
                    admin_only = excluded.admin_only,
                    hidden = excluded.hidden,
                    uploaded_by = excluded.uploaded_by,
                    uploaded_at = excluded.uploaded_at
            """, (filename, display_name, description, admin_only, hidden, current_username()))
            conn.commit()
            log_action(
                current_username(),
                f"{'оновив' if is_replace else 'завантажив новий'} шаблон документа: {filename}",
                details=f"тільки для адміністратора: {'так' if admin_only else 'ні'}"
            )
            flash(f"Шаблон «{filename}» успішно {'оновлено' if is_replace else 'завантажено'}", 'success')
        except Exception as e:
            logger.error(f"Помилка при завантаженні шаблону {filename}: {e}", exc_info=True)
            flash(f'Помилка при завантаженні шаблону: {e}', 'danger')
        finally:
            conn.close()

        return redirect(url_for('admin.manage_templates'))

    # -------------------- GET --------------------
    # Показуємо усі фізично наявні файли в template_word/, приєднуючи
    # метадані з БД, якщо вони є (файл міг бути покладений вручну, без
    # завантаження через цю сторінку - таким теж не повинно ламати список).
    rows = conn.execute("SELECT * FROM document_templates").fetchall()
    meta_by_filename = {row['filename']: dict(row) for row in rows}
    conn.close()

    templates = []
    if os.path.isdir(TEMPLATE_FOLDER):
        for f in sorted(os.listdir(TEMPLATE_FOLDER)):
            if not f.lower().endswith('.docx'):
                continue
            full_path = os.path.join(TEMPLATE_FOLDER, f)
            meta = meta_by_filename.get(f, {})
            templates.append({
                'filename': f,
                'display_name': meta.get('display_name') or f,
                'description': meta.get('description') or '',
                'admin_only': bool(meta.get('admin_only', 0)),
                'hidden': bool(meta.get('hidden', 0)),
                'uploaded_by': meta.get('uploaded_by') or '',
                'uploaded_at': meta.get('uploaded_at') or '',
                'size_kb': round(os.path.getsize(full_path) / 1024, 1),
            })

    return render_template('manage_templates.html', templates=templates)


@admin_bp.route('/admin/templates/<filename>/toggle_visibility', methods=['POST'])
@permission_required('manage_templates')
def toggle_template_visibility(filename):
    """
    Перемикає видимість шаблону в списках вибору при генерації
    (прихований <-> видимий), без потреби перезавантажувати файл.
    Для файлів, які лежать у template_word/ без запису в БД (додані
    вручну), запис створюється автоматично.
    """
    filename = secure_filename(filename)
    full_path = os.path.join(TEMPLATE_FOLDER, filename)

    if not os.path.isfile(full_path):
        flash(f"Файл шаблону «{filename}» не знайдено", 'danger')
        return redirect(url_for('admin.manage_templates'))

    conn = get_db()
    try:
        row = conn.execute("SELECT hidden FROM document_templates WHERE filename = ?", (filename,)).fetchone()
        if row is None:
            # Файл без метаданих (покладений вручну) - створюємо запис одразу прихованим
            conn.execute(
                "INSERT INTO document_templates (filename, display_name, hidden, uploaded_by) VALUES (?, ?, 1, ?)",
                (filename, filename, current_username())
            )
            new_hidden = 1
        else:
            new_hidden = 0 if row['hidden'] else 1
            conn.execute("UPDATE document_templates SET hidden = ? WHERE filename = ?", (new_hidden, filename))
        conn.commit()
        log_action(
            current_username(),
            f"{'приховав' if new_hidden else 'зробив видимим'} шаблон документа: {filename}"
        )
        flash(f"Шаблон «{filename}» тепер {'прихований зі' if new_hidden else 'видимий у'} списках вибору", 'success')
    except Exception as e:
        logger.error(f"Помилка при зміні видимості шаблону {filename}: {e}", exc_info=True)
        flash(f'Помилка при зміні видимості шаблону: {e}', 'danger')
    finally:
        conn.close()

    return redirect(url_for('admin.manage_templates'))


@admin_bp.route('/admin/templates/<filename>/download')
@permission_required('manage_templates')
def download_template(filename):
    """Віддає файл шаблону з template_word/ на завантаження (напр., щоб відредагувати його у Word і завантажити оновлену версію назад)."""
    filename = secure_filename(filename)
    full_path = os.path.join(TEMPLATE_FOLDER, filename)

    if not os.path.isfile(full_path):
        flash(f"Файл шаблону «{filename}» не знайдено", 'danger')
        return redirect(url_for('admin.manage_templates'))

    return send_file(full_path, as_attachment=True, download_name=filename)


@admin_bp.route('/admin/templates/<filename>/delete', methods=['POST'])
@permission_required('manage_templates')
def delete_template(filename):
    """Видаляє шаблон - сам файл із template_word/ та його метадані з БД."""
    filename = secure_filename(filename)
    full_path = os.path.join(TEMPLATE_FOLDER, filename)

    conn = get_db()
    try:
        conn.execute("DELETE FROM document_templates WHERE filename = ?", (filename,))
        conn.commit()
        if os.path.exists(full_path):
            os.remove(full_path)
        log_action(current_username(), f"видалив шаблон документа: {filename}")
        flash(f"Шаблон «{filename}» видалено", 'success')
    except Exception as e:
        logger.error(f"Помилка при видаленні шаблону {filename}: {e}", exc_info=True)
        flash(f'Помилка при видаленні шаблону: {e}', 'danger')
    finally:
        conn.close()

    return redirect(url_for('admin.manage_templates'))


@admin_bp.route('/admin/export_photos', methods=['GET', 'POST'])
@permission_required('group_export')
def export_photos():
    """
    Масове вивантаження фото студентів групи одним ZIP-архівом - для
    друку студентських квитків тощо. Можна забрати всіх студентів
    групи з фото одразу, або зняти позначку з окремих і завантажити
    лише вибраних. Кожен файл у архіві називається "Прізвище_Ім'я_По
    батькові.jpg" (по батькові пропускається, якщо не вказано) - готово
    вставляти в будь-яку програму верстки квитків без перейменування.
    """
    conn = get_db()
    conn.row_factory = sqlite3.Row

    if request.method == 'POST':
        group_id = request.form.get('group_id')
        student_ids = request.form.getlist('student_ids')
        if not student_ids:
            flash('Оберіть хоча б одного студента з фото', 'error')
            return redirect(url_for('admin.export_photos', group_id=group_id))

        placeholders = ','.join('?' for _ in student_ids)
        students = conn.execute(f"""
            SELECT id, last_name_UA, first_name_UA, middle_name_UA, photo
            FROM students WHERE id IN ({placeholders}) AND photo IS NOT NULL
        """, student_ids).fetchall()

        group = conn.execute(
            "SELECT name, degree_level, course FROM groups WHERE id=?", (group_id,)
        ).fetchone()

        # Паспортні дані (найновіший запис кожного студента) + ІПН -
        # для супровідного docx "Інформація для студентських квитків".
        id_card_rows = []
        if students:
            id_card_placeholders = ','.join('?' for _ in students)
            id_card_student_ids = [s['id'] for s in students]
            passport_by_student = {}
            for pd_row in conn.execute(f"""
                SELECT student_id, series, number, issued_by, issue_date, valid_until, unique_number
                FROM passport_documents
                WHERE student_id IN ({id_card_placeholders})
                ORDER BY id DESC
            """, id_card_student_ids).fetchall():
                passport_by_student.setdefault(pd_row['student_id'], pd_row)

            tax_id_by_student = {
                r['id']: r['tax_id'] for r in conn.execute(f"""
                    SELECT id, tax_id FROM students WHERE id IN ({id_card_placeholders})
                """, id_card_student_ids).fetchall()
            }

            sorted_students = sort_ukrainian(
                students,
                key_func=lambda s: f"{s['last_name_UA']} {s['first_name_UA']} {s['middle_name_UA']}"
            )
            for s in sorted_students:
                pd_row = passport_by_student.get(s['id'])
                name_parts = [s['last_name_UA'], s['first_name_UA'], s['middle_name_UA']]
                full_name = ' '.join(p.strip() for p in name_parts if p and p.strip())
                id_card_rows.append({
                    'full_name': full_name,
                    'passport': ((pd_row['series'] or '') + (pd_row['number'] or '')) if pd_row else '',
                    'issued_by': pd_row['issued_by'] if pd_row else '',
                    'issue_date': pd_row['issue_date'] if pd_row else '',
                    'valid_until': pd_row['valid_until'] if pd_row else '',
                    'unique_number': pd_row['unique_number'] if pd_row else '',
                    'tax_id': tax_id_by_student.get(s['id']) or '',
                })

        conn.close()

        if not students:
            flash('У жодного з обраних студентів немає завантаженого фото', 'error')
            return redirect(url_for('admin.export_photos', group_id=group_id))

        buffer = io.BytesIO()
        used_names = {}
        with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as zf:
            for s in students:
                photo_path = os.path.join('static', s['photo'])
                if not os.path.exists(photo_path):
                    continue
                ext = os.path.splitext(s['photo'])[1] or '.jpg'
                name_parts = [s['last_name_UA'], s['first_name_UA']]
                if s['middle_name_UA']:
                    name_parts.append(s['middle_name_UA'])
                base_name = '_'.join(p.strip() for p in name_parts if p and p.strip())
                # Про всяк випадок - якщо в групі раптом двоє тезок з
                # однаковим ПІБ, другий файл не повинен мовчки
                # перезаписати перший у архіві.
                arcname = f"{base_name}{ext}"
                if arcname in used_names:
                    used_names[arcname] += 1
                    arcname = f"{base_name}_{used_names[arcname]}{ext}"
                else:
                    used_names[arcname] = 1
                zf.write(photo_path, arcname=arcname)

            if id_card_rows:
                id_card_doc = build_id_card_info_docx(
                    id_card_rows,
                    group_name=(group['name'] if group else ''),
                    degree_level=(group['degree_level'] if group else ''),
                    course=(group['course'] if group else '')
                )
                zf.writestr('Інформація_для_студентських_квитків.docx', id_card_doc.read())

        buffer.seek(0)
        log_action(current_username(), f"вивантажив фото студентів (ZIP): {len(students)} шт., група ID {group_id}")
        safe_group_name = (group['name'] if group else 'group').replace(' ', '_')
        download_name = f"Фото_{safe_group_name}_{datetime.now().strftime('%Y-%m-%d')}.zip"
        return send_file(buffer, as_attachment=True, download_name=download_name, mimetype='application/zip')

    # ====================== GET ======================
    groups = conn.execute("""
        SELECT id, name, start_year, study_form, program_credits,
               name || ' (' || start_year || ', ' || study_form || ', ' || program_credits || ' кредитів)' AS display_name
        FROM groups WHERE archived = FALSE ORDER BY name COLLATE UKRAINIAN
    """).fetchall()

    selected_group_id = request.args.get('group_id', type=int)
    students = []
    if selected_group_id:
        students = conn.execute("""
            SELECT id, last_name_UA, first_name_UA, middle_name_UA, photo
            FROM students WHERE group_id = ? AND COALESCE(archived, 0) = 0
            ORDER BY last_name_UA COLLATE UKRAINIAN
        """, (selected_group_id,)).fetchall()
        students = sort_ukrainian(students, key_func=lambda s: f"{s['last_name_UA']} {s['first_name_UA']} {s['middle_name_UA']}")

    conn.close()
    return render_template(
        'export_photos.html',
        groups=groups, selected_group_id=selected_group_id, students=students,
    )


@admin_bp.route('/admin/group_export', methods=['GET', 'POST'])
@permission_required('group_export')
def group_export():
    """Сторінка вибору параметрів (група і/або рік народження, шаблон .docx) перед масовою генерацією документів студентів."""
    conn = get_db()
    conn.row_factory = sqlite3.Row

    groups = conn.execute("""
        SELECT id, name, start_year, study_form, program_credits,
               name || ' (' || start_year || ', ' || study_form || ', ' || program_credits || ' кредитів)' AS display_name
        FROM groups WHERE archived = FALSE ORDER BY name, start_year
    """).fetchall()

    available_templates = get_templates_with_metadata(is_admin=session.get('is_admin', False))
    default_template = available_templates[0]['path'] if available_templates else ''

    current_year = datetime.now().year
    years = list(range(1980, current_year + 1))
    students = []
    selected_group_id = request.args.get('group_id', type=int) if request.method == 'GET' else request.form.get('group_id', type=int)
    selected_year = request.args.get('birth_year', type=int) if request.method == 'GET' else request.form.get('birth_year', type=int)
    selected_template = request.args.get('template', default_template) if request.method == 'GET' else request.form.get('template', default_template)

    if selected_group_id:
        group_check = conn.execute("SELECT id FROM groups WHERE id=? AND archived=FALSE", (selected_group_id,)).fetchone()
        if not group_check:
            flash('Обрана група не існує або є архівною.', 'error')
            selected_group_id = None

    if request.method == 'POST':
        if not selected_group_id and not selected_year:
            flash('Будь ласка, оберіть групу або рік народження.', 'error')
        else:
            active_students = request.form.getlist('active_students')
            return redirect(url_for('admin.generate_group_docs', group_id=selected_group_id,
                                    birth_year=selected_year, template=selected_template,
                                    active_students=','.join(active_students)))

    if selected_group_id or selected_year:
        base_query = "SELECT * FROM students WHERE archived = FALSE"
        params = []
        if selected_group_id:
            base_query += " AND group_id=?"
            params.append(selected_group_id)
        if selected_year:
            base_query += " AND SUBSTR(birth_date, 7, 4) >= ?"
            params.append(str(selected_year))
        try:
            students = conn.execute(base_query, params).fetchall()
        except Exception as e:
            logger.error(f"Помилка при отриманні студентів: {e}")
            conn.close()
            return "Помилка бази даних", 500

    conn.close()
    return render_template('group_export.html', students=students, groups=groups, years=years,
                           selected_group_id=selected_group_id, selected_year=selected_year,
                           selected_template=selected_template,
                           available_templates=available_templates)


def _cleanup_old_jobs():
    now = time.time()
    with _JOBS_LOCK:
        stale = [jid for jid, j in _JOBS.items() if now - j['started_at'] > _JOB_TTL_SECONDS]
        for jid in stale:
            del _JOBS[jid]


def _run_generation_job(job_id, students_data, selected_template, batch_id, user_id, username, group_name, birth_year):
    """Виконується в окремому потоці: генерує документи по одному,
    оновлюючи прогрес у _JOBS[job_id] після кожного студента - навіть
    якщо один документ впаде з помилкою, решта продовжують генеруватись."""
    job = _JOBS[job_id]
    for student_dict, military_dict in students_data:
        with _JOBS_LOCK:
            job['current_name'] = f"{student_dict.get('last_name_UA','')} {student_dict.get('first_name_UA','')}"
        filename = f"{student_dict['last_name_UA']}_{student_dict['first_name_UA']}.docx".replace(" ", "_")
        full_path = os.path.join(office_editor.SESSIONS_DIR, f"{uuid.uuid4().hex}.docx")
        try:
            gen_doc(student_dict, military_dict, template=selected_template, out=full_path, user_name=username)
            doc_id = office_editor.create_editing_session(full_path, filename, user_id, batch_id=batch_id)
            with _JOBS_LOCK:
                job['succeeded'].append({
                    'doc_id': doc_id,
                    'name': f"{student_dict['last_name_UA']} {student_dict['first_name_UA']}",
                    'filename': filename,
                })
        except Exception as e:
            logger.error(f"Помилка при генерації документа для {student_dict.get('last_name_UA', '')}: {e}", exc_info=True)
            with _JOBS_LOCK:
                job['failed'].append({
                    'name': f"{student_dict.get('last_name_UA','')} {student_dict.get('first_name_UA','')}",
                    'error': str(e),
                })
        finally:
            with _JOBS_LOCK:
                job['done'] += 1

    with _JOBS_LOCK:
        job['complete'] = True
        job['current_name'] = ''

    log_action(
        username,
        f"масова генерація документів: {group_name}",
        details=f"шаблон: {selected_template}, рік нар.: {birth_year or 'всі'}, "
                f"успішно: {len(job['succeeded'])}, з помилкою: {len(job['failed'])}"
    )


@admin_bp.route('/admin/generate_group_docs', methods=['GET', 'POST'])
@permission_required('group_export')
def generate_group_docs():
    """Генерує .docx-документи (за обраним шаблоном) для всіх студентів, що підпадають під фільтр (група і/або рік народження) у фоновому потоці, показуючи прогрес-бар, а потім відкриває сторінку перегляду/редагування кожного в ONLYOFFICE перед завантаженням підсумкового ZIP-архіву (routes/office_editor.py)."""
    _cleanup_old_jobs()
    group_id = request.args.get('group_id', type=int) if request.method == 'GET' else request.form.get('group_id', type=int)
    birth_year = request.args.get('birth_year', type=int) if request.method == 'GET' else request.form.get('birth_year', type=int)
    selected_template = request.args.get('template', '') if request.method == 'GET' else request.form.get('template', '')
    active_students = request.args.get('active_students', '').split(',') if request.args.get('active_students') else []

    allowed_paths = {t['path'] for t in get_templates_with_metadata(is_admin=session.get('is_admin', False))}
    if selected_template not in allowed_paths:
        flash("У вас немає прав для генерації документів цим шаблоном", "danger")
        return redirect(url_for('admin.group_export'))

    if not group_id and not birth_year:
        flash('Оберіть групу або рік народження для генерації документів.', 'error')
        return redirect(url_for('admin.group_export'))

    conn = get_db()
    conn.row_factory = sqlite3.Row
    base_query = """
        SELECT s.*,
               g.name || ' (' || g.start_year || ', ' || g.study_form || ', ' || g.program_credits || ' кредитів)' AS group_name,
               g.study_form, g.start_year, g.program_credits,
               g.qualification_name, g.degree_level, g.specialty, g.educational_program, g.knowledge_area,
               g.qualification_name_en, g.degree_level_en, g.specialty_en, g.isced_f_ua, g.isced_f_en, g.educational_program_en, g.knowledge_area_en,
               il.name_ua AS institution_name_and_status, il.name_en AS institution_name_and_status_en,
               il.short_name_ua AS license_short_name_ua, il.short_name_en AS license_short_name_en,
               g.entry_requirements, g.entry_requirements_en,
               g.learning_outcomes, g.learning_outcomes_en, g.program_includes, g.program_includes_en,
               g.entry_requirements_reduced, g.entry_requirements_reduced_en,
               g.learning_outcomes_reduced, g.learning_outcomes_reduced_en,
               g.program_includes_reduced, g.program_includes_reduced_en
        FROM students s LEFT JOIN groups g ON s.group_id = g.id
                         LEFT JOIN institution_licenses il ON s.license_id = il.id
        WHERE s.archived = FALSE
    """
    params = []
    if group_id:
        base_query += " AND s.group_id=?"
        params.append(group_id)
    if birth_year:
        base_query += " AND SUBSTR(s.birth_date, 7, 4) >= ?"
        params.append(str(birth_year))

    try:
        students = conn.execute(base_query, params).fetchall()
        if not students:
            conn.close()
            return "Студенты не найдены по заданным фильтрам", 404
    except Exception as e:
        logger.error(f"Ошибка при выполнении SQL-запроса: {e}")
        conn.close()
        return "Ошибка базы данных", 500

    if active_students and active_students[0]:
        students = [s for s in students if str(s['id']) in active_students]

    group_name = "Зі всіх груп"
    if group_id and students:
        group_name = students[0]['group_name'] if students[0]['group_name'] else f"Група_{group_id}"

    batch_id = office_editor.new_batch_id()

    # Забираємо всі дані студентів (і військові дані) заздалегідь, поки
    # з'єднання з базою відкрите в цьому запиті - фоновий потік більше
    # не звертатиметься до conn з цього обробника (SQLite-з'єднання
    # прив'язане до потоку, у якому було створене).
    students_data = []
    for student in students:
        student_dict = dict(student)
        military = conn.execute("SELECT * FROM military WHERE student_id=?", (student['id'],)).fetchone()
        students_data.append((student_dict, dict(military) if military else {}))
    conn.close()

    job_id = uuid.uuid4().hex
    with _JOBS_LOCK:
        _JOBS[job_id] = {
            'total': len(students_data),
            'done': 0,
            'current_name': '',
            'succeeded': [],
            'failed': [],
            'complete': False,
            'batch_id': batch_id,
            'group_name': group_name,
            'started_at': time.time(),
        }

    thread = threading.Thread(
        target=_run_generation_job,
        args=(job_id, students_data, selected_template, batch_id, session['user_id'],
              current_username(), group_name, birth_year),
        daemon=True,
    )
    thread.start()

    return render_template('group_generate_progress.html', job_id=job_id, group_name=group_name, total=len(students_data))


@admin_bp.route('/admin/generate_group_docs/status/<job_id>')
@permission_required('group_export')
def generate_group_docs_status(job_id):
    """JSON-статус фонової генерації - опитується сторінкою прогресу через AJAX."""
    with _JOBS_LOCK:
        job = _JOBS.get(job_id)
        if not job:
            return {'error': 'not_found'}, 404
        return {
            'total': job['total'],
            'done': job['done'],
            'current_name': job['current_name'],
            'complete': job['complete'],
            'succeeded_count': len(job['succeeded']),
            'failed_count': len(job['failed']),
        }


@admin_bp.route('/admin/generate_group_docs/result/<job_id>')
@permission_required('group_export')
def generate_group_docs_result(job_id):
    """Показує підсумок завершеної фонової генерації: перелік готових документів (з переходом у ONLYOFFICE) і, за наявності, список тих, що не вдалося згенерувати."""
    with _JOBS_LOCK:
        job = _JOBS.get(job_id)
    if not job or not job['complete']:
        flash('Завдання генерації не знайдено або ще не завершено', 'warning')
        return redirect(url_for('admin.group_export'))

    if not job['succeeded']:
        flash('Не вдалося згенерувати жодного документа', 'danger')
        return redirect(url_for('admin.group_export'))

    return render_template(
        'group_docs_preview.html',
        items=job['succeeded'],
        failed=job['failed'],
        batch_id=job['batch_id'],
        group_name=job['group_name'],
    )