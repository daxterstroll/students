"""
routes/admin/documents.py
==========================
Документи про освіту й паспортні дані студентів (+ скани-вкладення),
масовий імпорт документів про освіту (двоетапний: попередній перегляд
-> підтвердження) і паспортних даних (одноетапний), разом із
допоміжними функціями розпізнавання Excel-комірок для імпорту.
"""
from flask import render_template, request, redirect, url_for, flash, session, send_file
from routes.db import get_db
from routes.utils import log_action, permission_required, logger
from routes.helpers import current_username, sort_ukrainian
from routes.admin import admin_bp
import sqlite3
import os
from routes.utils import get_attachments, save_multiple_attachments
from routes.admin import allowed_file
from werkzeug.utils import secure_filename
import pandas as pd
import json
import uuid
import re
import time
from openpyxl import load_workbook
from rapidfuzz import process, fuzz
from deep_translator import GoogleTranslator

ALLOWED_EXTENSIONS = {'xlsx'}

TEMP_PREVIEW_FOLDER = "temp_preview"

translator = GoogleTranslator(source="auto", target="en")
translation_cache = {}


@admin_bp.route('/admin/manage_education_documents', methods=['GET', 'POST'])
@permission_required('manage_education_documents')
def manage_education_documents():
    """CRUD-сторінка документів про освіту студентів (атестат/диплом попереднього рівня) разом із додатковими даними про визнання іноземного документа (foreign_education_docs)."""
    db = get_db()
    cursor = db.cursor()

    cursor.execute("""
        SELECT id, last_name_UA, first_name_UA FROM students
        WHERE archived = FALSE ORDER BY last_name_UA, first_name_UA
    """)
    students = cursor.fetchall()

    cursor.execute("SELECT id, name FROM groups WHERE archived = FALSE ORDER BY name")
    groups = cursor.fetchall()

    # --- Отримуємо студента, якщо перейшли з картки ---
    student = None
    student_id_param = request.args.get('student_id', type=int)
    if student_id_param:
        cursor.execute("""
            SELECT id, last_name_UA, first_name_UA, middle_name_UA, group_id
            FROM students WHERE id = ?
        """, (student_id_param,))
        srow = cursor.fetchone()
        if srow:
            student = dict(srow)
            cursor.execute("SELECT name FROM groups WHERE id = ?", (student['group_id'],))
            grow = cursor.fetchone()
            student['group_name'] = grow['name'] if grow else ''

    selected_group_id = request.args.get('group_id', type=int)
    if not selected_group_id and student:
        selected_group_id = student['group_id']

    students_without_docs = []
    if selected_group_id:
        cursor.execute("""
            SELECT s.id, s.last_name_UA, s.first_name_UA
            FROM students s
            WHERE s.group_id = ? AND s.archived = FALSE
              AND s.id NOT IN (SELECT student_id FROM education_documents)
            ORDER BY s.last_name_UA, s.first_name_UA
        """, (selected_group_id,))
        students_without_docs = cursor.fetchall()

    # Основний запит документів
    cursor.execute("""
        SELECT g.id AS group_id, g.name AS group_name, s.id AS student_id,
               s.last_name_UA, s.first_name_UA,
               ed.id AS doc_id, ed.document_type, ed.document_type_en, ed.document_number,
               ed.institution_name, ed.institution_name_en, ed.country, ed.country_en, ed.completion_date,
               fed.reference_number, fed.reference_institution, fed.reference_institution_en,
               fed.reference_country, fed.reference_country_en, fed.reference_issue_date,
               fed.recognition_certificate_number, fed.recognition_issuer,
               fed.recognition_issuer_en, fed.recognition_date
        FROM education_documents ed
        INNER JOIN students s ON ed.student_id = s.id
        INNER JOIN groups g ON s.group_id = g.id
        LEFT JOIN foreign_education_docs fed ON ed.id = fed.education_doc_id
        WHERE s.archived = FALSE AND g.archived = FALSE
        ORDER BY g.name, s.last_name_UA, s.first_name_UA, ed.id
    """)
    rows = cursor.fetchall()

    documents_by_group = {}
    for row in rows:
        gid = row['group_id']
        if gid not in documents_by_group:
            documents_by_group[gid] = {'group_name': row['group_name'], 'docs': []}
        doc = dict(row)
        doc['attachments'] = get_attachments(db, 'education_document', doc['doc_id'])
        documents_by_group[gid]['docs'].append(doc)

    sorted_documents_by_group = sorted(documents_by_group.items(), key=lambda x: x[1]['group_name'])

    # Перевіряємо doc_id для студента
    student_doc_id = None
    if student:
        for gid, gdata in documents_by_group.items():
            for doc in gdata['docs']:
                if doc['student_id'] == student['id']:
                    student_doc_id = doc['doc_id']
                    break
            if student_doc_id:
                break

    # ====================== POST ======================
    if request.method == 'POST':
        action = request.form.get('action')

        if action == 'delete':
            doc_id = request.form.get('doc_id')
            try:
                # Спершу самі файли сканів з диска, а тоді записи -
                # інакше видалення документа лишало б "осиротілі" файли.
                for att in get_attachments(db, 'education_document', doc_id):
                    att_path = os.path.join('static', att['file_path'])
                    if os.path.exists(att_path):
                        os.remove(att_path)
                cursor.execute("DELETE FROM attachments WHERE entity_type='education_document' AND entity_id=?", (doc_id,))
                cursor.execute("DELETE FROM foreign_education_docs WHERE education_doc_id = ?", (doc_id,))
                cursor.execute("DELETE FROM education_documents WHERE id = ?", (doc_id,))
                db.commit()
                log_action(current_username(), f"ВИДАЛИВ документ про освіту ID {doc_id}")
                flash('Документ успішно видалено', 'success')
            except sqlite3.Error as e:
                db.rollback()
                logger.error(f"Помилка видалення документа про освіту (ID {doc_id}): {e}", exc_info=True)
                flash(f'Помилка видалення: {e}', 'danger')

        elif action == 'delete_attachment':
            attachment_id = request.form.get('attachment_id')
            try:
                att = cursor.execute("SELECT file_path FROM attachments WHERE id=?", (attachment_id,)).fetchone()
                if att:
                    att_path = os.path.join('static', att['file_path'])
                    if os.path.exists(att_path):
                        os.remove(att_path)
                    cursor.execute("DELETE FROM attachments WHERE id=?", (attachment_id,))
                    db.commit()
                    flash('Скан видалено', 'success')
                else:
                    flash('Файл не знайдено', 'error')
            except Exception as e:
                db.rollback()
                logger.error(f"Помилка видалення вкладення (ID {attachment_id}): {e}", exc_info=True)
                flash(f'Помилка видалення файлу: {e}', 'danger')
            return redirect(url_for('admin.manage_education_documents', group_id=selected_group_id))

        elif action == 'edit':
            doc_id = request.form.get('doc_id')
            student_id = request.form.get('student_id')

            country = request.form.get('country') or ''
            country_en = request.form.get('country_en') or ''

            try:
                cursor.execute("SELECT student_id FROM education_documents WHERE id = ?", (doc_id,))
                existing = cursor.fetchone()
                if not existing:
                    flash('Документ не знайдено', 'danger')
                    return redirect(url_for('admin.manage_education_documents', group_id=selected_group_id))

                if student_id:
                    cursor.execute("SELECT id FROM students WHERE id = ? AND archived = FALSE", (student_id,))
                    if not cursor.fetchone():
                        flash('Обраний студент не існує або заархівований', 'danger')
                        return redirect(url_for('admin.manage_education_documents', group_id=selected_group_id))
                else:
                    student_id = existing[0]

                # Оновлення основного документа (зберігаємо country як ввели)
                cursor.execute("""
                    UPDATE education_documents SET
                        student_id=?, document_type=?, document_type_en=?, document_number=?,
                        institution_name=?, institution_name_en=?, country=?, country_en=?, completion_date=?
                    WHERE id=?
                """, (student_id, request.form.get('document_type'), request.form.get('document_type_en'),
                      request.form.get('document_number'), request.form.get('institution_name'),
                      request.form.get('institution_name_en'), country, country_en,
                      request.form.get('completion_date'), doc_id))

                # Foreign fields
                reference_number = request.form.get('reference_number') or None
                reference_institution = request.form.get('reference_institution') or None
                reference_institution_en = request.form.get('reference_institution_en') or None
                reference_country = request.form.get('reference_country') or None
                reference_country_en = request.form.get('reference_country_en') or None
                reference_issue_date = request.form.get('reference_issue_date') or None
                recognition_certificate_number = request.form.get('recognition_certificate_number') or None
                recognition_issuer = request.form.get('recognition_issuer') or None
                recognition_issuer_en = request.form.get('recognition_issuer_en') or None
                recognition_date = request.form.get('recognition_date') or None

                # Рішення приймаємо тільки на основі того, чи заповнені самі поля довідки/визнання,
                # а НЕ на основі значення country. Це прибирає випадкове видалення довідки
                # через неправильно передане/незмінене поле "Країна".
                has_foreign_data = any([
                    reference_number, reference_institution, reference_country,
                    reference_issue_date, recognition_certificate_number,
                    recognition_issuer, recognition_date
                ])

                cursor.execute("SELECT id FROM foreign_education_docs WHERE education_doc_id = ?", (doc_id,))
                foreign_exists = cursor.fetchone()

                if foreign_exists:
                    if has_foreign_data:
                        cursor.execute("""
                            UPDATE foreign_education_docs SET
                                reference_number=?, reference_institution=?, reference_institution_en=?,
                                reference_country=?, reference_country_en=?, reference_issue_date=?,
                                recognition_certificate_number=?, recognition_issuer=?,
                                recognition_issuer_en=?, recognition_date=?
                            WHERE education_doc_id=?
                        """, (reference_number, reference_institution, reference_institution_en,
                              reference_country, reference_country_en, reference_issue_date,
                              recognition_certificate_number, recognition_issuer,
                              recognition_issuer_en, recognition_date, doc_id))
                    else:
                        # Видаляємо запис тільки якщо користувач сам очистив усі поля довідки/визнання
                        cursor.execute("DELETE FROM foreign_education_docs WHERE education_doc_id = ?", (doc_id,))
                else:
                    if has_foreign_data:
                        cursor.execute("""
                            INSERT INTO foreign_education_docs (
                                education_doc_id, reference_number, reference_institution, reference_institution_en,
                                reference_country, reference_country_en, reference_issue_date,
                                recognition_certificate_number, recognition_issuer, recognition_issuer_en, recognition_date
                            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                        """, (doc_id, reference_number, reference_institution, reference_institution_en,
                              reference_country, reference_country_en, reference_issue_date,
                              recognition_certificate_number, recognition_issuer, recognition_issuer_en, recognition_date))

                db.commit()

                new_scans = [f for f in request.files.getlist('document_scans') if f and f.filename]
                if new_scans:
                    save_multiple_attachments(db, 'education_document', doc_id, new_scans, 'education_documents', current_username())
                    db.commit()

                log_action(
                    current_username(),
                    f"редагував документ про освіту ID {doc_id}",
                    details=f"країна: {country}"
                )
                flash('Документ успішно оновлено', 'success')

            except sqlite3.Error as e:
                db.rollback()
                logger.error(f"Помилка БД при редагуванні документа про освіту (ID {doc_id}): {e}", exc_info=True)
                flash(f'Помилка при редагуванні: {e}', 'danger')
            except Exception as e:
                db.rollback()
                logger.error(f"Непередбачена помилка при редагуванні документа про освіту (ID {doc_id}): {e}", exc_info=True)
                flash(f'Непередбачена помилка: {str(e)}', 'danger')

            return redirect(url_for('admin.manage_education_documents', group_id=selected_group_id))

        # === Додавання нового документа ===
        else:
            country = request.form.get('country') or ''
            country_en = request.form.get('country_en') or ''

            student_id = request.form.get('student_id')
            document_type = request.form.get('document_type')
            document_type_en = request.form.get('document_type_en')
            document_number = request.form.get('document_number')
            institution_name = request.form.get('institution_name')
            institution_name_en = request.form.get('institution_name_en')
            completion_date = request.form.get('completion_date')

            reference_number = request.form.get('reference_number') or None
            reference_institution = request.form.get('reference_institution') or None
            reference_institution_en = request.form.get('reference_institution_en') or None
            reference_country = request.form.get('reference_country') or None
            reference_country_en = request.form.get('reference_country_en') or None
            reference_issue_date = request.form.get('reference_issue_date') or None
            recognition_certificate_number = request.form.get('recognition_certificate_number') or None
            recognition_issuer = request.form.get('recognition_issuer') or None
            recognition_issuer_en = request.form.get('recognition_issuer_en') or None
            recognition_date = request.form.get('recognition_date') or None

            try:
                if not student_id:
                    raise ValueError("Не обрано студента")

                cursor.execute("""
                    INSERT INTO education_documents (
                        student_id, document_type, document_type_en, document_number,
                        institution_name, institution_name_en, country, country_en, completion_date
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (student_id, document_type, document_type_en, document_number,
                      institution_name, institution_name_en, country, country_en, completion_date))

                education_doc_id = cursor.lastrowid

                has_foreign_data = any([
                    reference_number, reference_institution, reference_country,
                    reference_issue_date, recognition_certificate_number,
                    recognition_issuer, recognition_date
                ])

                if has_foreign_data:
                    cursor.execute("""
                        INSERT INTO foreign_education_docs (
                            education_doc_id, reference_number, reference_institution, reference_institution_en,
                            reference_country, reference_country_en, reference_issue_date,
                            recognition_certificate_number, recognition_issuer, recognition_issuer_en, recognition_date
                        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """, (education_doc_id, reference_number, reference_institution, reference_institution_en,
                          reference_country, reference_country_en, reference_issue_date,
                          recognition_certificate_number, recognition_issuer, recognition_issuer_en, recognition_date))

                db.commit()
                log_action(
                    current_username(),
                    f"додав документ про освіту: {document_type} №{document_number}",
                    details=f"країна: {country}"
                )

                new_scans = [f for f in request.files.getlist('document_scans') if f and f.filename]
                if new_scans:
                    save_multiple_attachments(db, 'education_document', education_doc_id, new_scans, 'education_documents', current_username())
                    db.commit()

                flash('Документ успішно додано', 'success')

            except (sqlite3.Error, ValueError) as e:
                db.rollback()
                flash(f'Помилка при додаванні: {str(e)}', 'danger')

        return redirect(url_for('admin.manage_education_documents', group_id=selected_group_id))

    cursor.close()
    return render_template(
        'manage_education_documents.html',
        groups=groups,
        selected_group_id=selected_group_id,
        students_without_docs=students_without_docs,
        documents_by_group=sorted_documents_by_group,
        students=students,
        student=student,
        student_doc_id=student_doc_id
    )


@admin_bp.route('/admin/manage_passport_documents', methods=['GET', 'POST'])
@permission_required('manage_passport_documents')
def manage_passport_documents():
    """CRUD-сторінка паспортних даних студентів (паспорт-книжка або
    ID-картка) - за тим самим зразком, що й документи про освіту, але
    без підтаблиці "іноземний документ" і з лімітом 5 сканів на запис."""
    db = get_db()
    cursor = db.cursor()
    MAX_FILES = 5

    cursor.execute("""
        SELECT id, last_name_UA, first_name_UA FROM students
        WHERE archived = FALSE ORDER BY last_name_UA, first_name_UA
    """)
    students = cursor.fetchall()

    cursor.execute("SELECT id, name FROM groups WHERE archived = FALSE ORDER BY name")
    groups = cursor.fetchall()

    student = None
    student_id_param = request.args.get('student_id', type=int)
    if student_id_param:
        cursor.execute("""
            SELECT id, last_name_UA, first_name_UA, middle_name_UA, group_id
            FROM students WHERE id = ?
        """, (student_id_param,))
        srow = cursor.fetchone()
        if srow:
            student = dict(srow)
            cursor.execute("SELECT name FROM groups WHERE id = ?", (student['group_id'],))
            grow = cursor.fetchone()
            student['group_name'] = grow['name'] if grow else ''

    selected_group_id = request.args.get('group_id', type=int)
    if not selected_group_id and student:
        selected_group_id = student['group_id']

    students_without_docs = []
    if selected_group_id:
        cursor.execute("""
            SELECT s.id, s.last_name_UA, s.first_name_UA
            FROM students s
            WHERE s.group_id = ? AND s.archived = FALSE
              AND s.id NOT IN (SELECT student_id FROM passport_documents)
            ORDER BY s.last_name_UA, s.first_name_UA
        """, (selected_group_id,))
        students_without_docs = cursor.fetchall()

    cursor.execute("""
        SELECT g.id AS group_id, g.name AS group_name, s.id AS student_id,
               s.last_name_UA, s.first_name_UA,
               pd.id AS doc_id, pd.document_type, pd.series, pd.number,
               pd.issued_by, pd.issue_date, pd.valid_until, pd.unique_number
        FROM passport_documents pd
        INNER JOIN students s ON pd.student_id = s.id
        INNER JOIN groups g ON s.group_id = g.id
        WHERE s.archived = FALSE AND g.archived = FALSE
        ORDER BY g.name, s.last_name_UA, s.first_name_UA, pd.id
    """)
    rows = cursor.fetchall()

    documents_by_group = {}
    for row in rows:
        gid = row['group_id']
        if gid not in documents_by_group:
            documents_by_group[gid] = {'group_name': row['group_name'], 'docs': []}
        doc = dict(row)
        doc['attachments'] = get_attachments(db, 'passport_document', doc['doc_id'])
        documents_by_group[gid]['docs'].append(doc)

    sorted_documents_by_group = sorted(documents_by_group.items(), key=lambda x: x[1]['group_name'])

    student_doc_id = None
    if student:
        for gid, gdata in documents_by_group.items():
            for doc in gdata['docs']:
                if doc['student_id'] == student['id']:
                    student_doc_id = doc['doc_id']
                    break
            if student_doc_id:
                break

    # ====================== POST ======================
    if request.method == 'POST':
        action = request.form.get('action')

        if action == 'delete':
            doc_id = request.form.get('doc_id')
            try:
                for att in get_attachments(db, 'passport_document', doc_id):
                    att_path = os.path.join('static', att['file_path'])
                    if os.path.exists(att_path):
                        os.remove(att_path)
                cursor.execute("DELETE FROM attachments WHERE entity_type='passport_document' AND entity_id=?", (doc_id,))
                cursor.execute("DELETE FROM passport_documents WHERE id = ?", (doc_id,))
                db.commit()
                log_action(current_username(), f"ВИДАЛИВ паспортні дані ID {doc_id}")
                flash('Запис успішно видалено', 'success')
            except sqlite3.Error as e:
                db.rollback()
                logger.error(f"Помилка видалення паспортних даних (ID {doc_id}): {e}", exc_info=True)
                flash(f'Помилка видалення: {e}', 'danger')

        elif action == 'delete_attachment':
            attachment_id = request.form.get('attachment_id')
            try:
                att = cursor.execute("SELECT file_path FROM attachments WHERE id=?", (attachment_id,)).fetchone()
                if att:
                    att_path = os.path.join('static', att['file_path'])
                    if os.path.exists(att_path):
                        os.remove(att_path)
                    cursor.execute("DELETE FROM attachments WHERE id=?", (attachment_id,))
                    db.commit()
                    flash('Скан видалено', 'success')
                else:
                    flash('Файл не знайдено', 'error')
            except Exception as e:
                db.rollback()
                logger.error(f"Помилка видалення вкладення (ID {attachment_id}): {e}", exc_info=True)
                flash(f'Помилка видалення файлу: {e}', 'danger')
            return redirect(url_for('admin.manage_passport_documents', group_id=selected_group_id))

        elif action == 'edit':
            doc_id = request.form.get('doc_id')
            student_id = request.form.get('student_id')
            document_type = request.form.get('document_type')

            try:
                existing = cursor.execute("SELECT student_id FROM passport_documents WHERE id = ?", (doc_id,)).fetchone()
                if not existing:
                    flash('Запис не знайдено', 'danger')
                    return redirect(url_for('admin.manage_passport_documents', group_id=selected_group_id))

                if student_id:
                    if not cursor.execute("SELECT id FROM students WHERE id = ? AND archived = FALSE", (student_id,)).fetchone():
                        flash('Обраний студент не існує або заархівований', 'danger')
                        return redirect(url_for('admin.manage_passport_documents', group_id=selected_group_id))
                else:
                    student_id = existing[0]

                if document_type not in ('Паспорт (книжка)', 'ID-картка'):
                    raise ValueError("Невірний тип документа")

                existing_count = cursor.execute(
                    "SELECT COUNT(*) FROM attachments WHERE entity_type='passport_document' AND entity_id=?", (doc_id,)
                ).fetchone()[0]
                new_scans = [f for f in request.files.getlist('document_scans') if f and f.filename]
                if existing_count + len(new_scans) > MAX_FILES:
                    flash(f'Забагато файлів - максимум {MAX_FILES} на запис (вже є {existing_count}).', 'error')
                    return redirect(url_for('admin.manage_passport_documents', group_id=selected_group_id))

                cursor.execute("""
                    UPDATE passport_documents SET
                        student_id=?, document_type=?, series=?, number=?,
                        issued_by=?, issue_date=?, valid_until=?, unique_number=?
                    WHERE id=?
                """, (
                    student_id, document_type, request.form.get('series') or None, request.form.get('number'),
                    request.form.get('issued_by') or None, request.form.get('issue_date') or None,
                    request.form.get('valid_until') or None, request.form.get('unique_number') or None,
                    doc_id,
                ))
                db.commit()

                if new_scans:
                    save_multiple_attachments(db, 'passport_document', doc_id, new_scans, 'passport_documents', current_username())
                    db.commit()

                log_action(current_username(), f"редагував паспортні дані ID {doc_id}", details=document_type)
                flash('Дані успішно оновлено', 'success')

            except (sqlite3.Error, ValueError) as e:
                db.rollback()
                logger.error(f"Помилка редагування паспортних даних (ID {doc_id}): {e}", exc_info=True)
                flash(f'Помилка при редагуванні: {e}', 'danger')

            return redirect(url_for('admin.manage_passport_documents', group_id=selected_group_id))

        # === Додавання нового запису ===
        else:
            student_id = request.form.get('student_id')
            document_type = request.form.get('document_type')

            try:
                if not student_id:
                    raise ValueError("Не обрано студента")
                if document_type not in ('Паспорт (книжка)', 'ID-картка'):
                    raise ValueError("Невірний тип документа")

                new_scans = [f for f in request.files.getlist('document_scans') if f and f.filename]
                if len(new_scans) > MAX_FILES:
                    raise ValueError(f"Забагато файлів - максимум {MAX_FILES} на запис")

                cursor.execute("""
                    INSERT INTO passport_documents (
                        student_id, document_type, series, number, issued_by, issue_date, valid_until, unique_number
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """, (
                    student_id, document_type, request.form.get('series') or None, request.form.get('number'),
                    request.form.get('issued_by') or None, request.form.get('issue_date') or None,
                    request.form.get('valid_until') or None, request.form.get('unique_number') or None,
                ))
                passport_doc_id = cursor.lastrowid
                db.commit()

                if new_scans:
                    save_multiple_attachments(db, 'passport_document', passport_doc_id, new_scans, 'passport_documents', current_username())
                    db.commit()

                log_action(current_username(), f"додав паспортні дані: {document_type} №{request.form.get('number')}")
                flash('Запис успішно додано', 'success')

            except (sqlite3.Error, ValueError) as e:
                db.rollback()
                flash(f'Помилка при додаванні: {str(e)}', 'danger')

        return redirect(url_for('admin.manage_passport_documents', group_id=selected_group_id))

    cursor.close()
    return render_template(
        'manage_passport_documents.html',
        groups=groups,
        selected_group_id=selected_group_id,
        students_without_docs=students_without_docs,
        documents_by_group=sorted_documents_by_group,
        students=students,
        student=student,
        student_doc_id=student_doc_id,
        max_files=MAX_FILES,
    )


@admin_bp.route('/admin/import_passport_documents', methods=['GET', 'POST'])
@permission_required('import_passport_documents')
def import_passport_documents():
    """
    Масовий імпорт паспортних даних з Excel: кожен рядок - один
    студент (знаходиться нечітким пошуком за ПІБ серед усіх активних
    студентів, routes/admin.fuzzy_find_student). Якщо студент уже має
    паспортні дані - додає ще один запис (історія), не перезаписує.
    """
    if request.method == 'POST':
        file = request.files.get('excel_file')
        if not file or not allowed_file(file.filename):
            flash("⚠️ Оберіть коректний файл .xlsx", "error")
            return redirect(url_for('admin.import_passport_documents'))

        filename = f"passport_import_{int(time.time())}.xlsx"
        filepath = os.path.join('static', 'uploads', 'tmp', filename)
        os.makedirs(os.path.dirname(filepath), exist_ok=True)
        file.save(filepath)

        conn = get_db()
        cursor = conn.cursor()
        inserted, skipped = 0, 0
        try:
            wb = load_workbook(filepath)
            sheet = wb.active
            for i, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                if not row or all(cell is None for cell in row):
                    continue
                try:
                    fio = row[0]
                    document_type = str(row[1]).strip() if len(row) > 1 and row[1] else None
                    number = str(row[3]).strip() if len(row) > 3 and row[3] not in (None, '') else None

                    if not fio or not document_type or not number:
                        flash(f"❗ Рядок {i}: не заповнено ПІБ, тип документа чи номер - пропущено")
                        skipped += 1
                        continue
                    if document_type not in ('Паспорт (книжка)', 'ID-картка'):
                        flash(f"❗ Рядок {i}: невідомий тип документа '{document_type}' (очікується 'Паспорт (книжка)' або 'ID-картка') - пропущено")
                        skipped += 1
                        continue

                    student_id, matched_name, score = fuzzy_find_student(cursor, str(fio).strip())
                    if not student_id:
                        flash(f"❗ Рядок {i}: студента не знайдено за ПІБ '{fio}' - пропущено")
                        skipped += 1
                        continue

                    series = str(row[2]).strip() if len(row) > 2 and row[2] not in (None, '') else None
                    issued_by = str(row[4]).strip() if len(row) > 4 and row[4] not in (None, '') else None
                    issue_date = str(row[5]).strip().replace('-', '.') if len(row) > 5 and row[5] not in (None, '') else None
                    valid_until = str(row[6]).strip().replace('-', '.') if len(row) > 6 and row[6] not in (None, '') else None
                    unique_number = str(row[7]).strip() if len(row) > 7 and row[7] not in (None, '') else None

                    cursor.execute("""
                        INSERT INTO passport_documents (
                            student_id, document_type, series, number, issued_by, issue_date, valid_until, unique_number
                        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                    """, (student_id, document_type, series, number, issued_by, issue_date, valid_until, unique_number))
                    inserted += 1
                except Exception as e:
                    logger.error(f"Row {i} error (passport import): {e}")
                    flash(f"❗ Рядок {i}: помилка обробки - {e}")
                    skipped += 1
                    continue

            conn.commit()
            log_action(
                current_username(),
                f"імпорт паспортних даних з Excel: додано {inserted}, пропущено {skipped}",
                details=f"файл: {filename}"
            )
            flash(f"✅ Імпорт завершено. Додано: {inserted}, пропущено: {skipped}", "success")
        except Exception as e:
            conn.rollback()
            flash(f"⚠️ Помилка при імпорті Excel: {e}", "error")
            logger.error(f"Error importing passport Excel: {e}")
        finally:
            if os.path.exists(filepath):
                os.remove(filepath)
            conn.close()

        return redirect(url_for('admin.import_passport_documents'))

    return render_template('import_passport_documents.html')


def save_preview_to_file(preview):
    """Зберігає проміжний результат парсингу Excel-файлу імпорту документів про освіту у тимчасовий JSON-файл (щоб не тримати великі дані в сесії) і повертає його ID."""
    os.makedirs(TEMP_PREVIEW_FOLDER, exist_ok=True)
    preview_id = str(uuid.uuid4())
    path = os.path.join(TEMP_PREVIEW_FOLDER, f"{preview_id}.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(preview, f, ensure_ascii=False)
    return preview_id


def load_preview_from_file(preview_id):
    """Завантажує раніше збережений (save_preview_to_file) прев'ю-результат імпорту за його ID."""
    path = os.path.join(TEMP_PREVIEW_FOLDER, f"{preview_id}.json")
    if not os.path.exists(path):
        return []
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def fuzzy_find_student(cursor, full_name, threshold=80):
    """Нечіткий пошук студента за повним ім'ям (ПІБ) серед активних студентів, використовуючи rapidfuzz; повертає ID найкращого збігу, якщо він проходить поріг схожості threshold."""
    cursor.execute("SELECT id, last_name_UA, first_name_UA, middle_name_UA FROM students WHERE archived=FALSE")
    students = cursor.fetchall()
    names = ["{} {} {}".format(s["last_name_UA"], s["first_name_UA"], s["middle_name_UA"]) for s in students]
    matches = process.extract(full_name, names, scorer=fuzz.ratio, limit=3)
    for match_name, score, idx in matches:
        if score >= threshold:
            return students[idx]["id"], match_name, score
    return None, None, None


def translate_to_en(text):
    """Перекладає текст на англійську через GoogleTranslator з кешуванням результатів (translation_cache), щоб не робити повторні запити до перекладача для однакового тексту."""
    if not text:
        return ""
    if text in translation_cache:
        return translation_cache[text]
    try:
        result = translator.translate(text)
        translation_cache[text] = result
        return result
    except Exception as e:
        logger.warning(f"Помилка перекладу '{text}': {e}")
        return text


def find_country(text):
    """Визначає країну (укр./англ. назву) за ключовими словами в тексті установи/довідки; за замовчуванням вважає, що це Україна."""
    countries = {
        "польща": "Poland",
        "poland": "Poland",
        "німеччина": "Germany",
        "germany": "Germany",
        "чех": "Czech Republic",
        "czech": "Czech Republic",
        "словач": "Slovakia",
        "slovakia": "Slovakia"
    }

    lower = text.lower()

    for key, en in countries.items():
        if key in lower:
            return key.capitalize(), en

    if "україна " in lower:
        return "Україна ", "Ukraine"

    return "Україна", "Ukraine"


def parse_document(text: str):
    """Розбирає текстовий опис документа про освіту (формат 'Тип документа Номер; ДД.ММ.РРРР; Ким видано: Установа') на структуровані поля (тип, номер, дата, установа, країна) для імпорту."""
    text = text.strip()
    pattern = re.compile(
        r'^(?P<type>[^;]+?)\s*;\s*(?P<date>\d{2}\.\d{2}\.\d{4})\s*;\s*Ким видано:\s*(?P<institution>.+?)$',
        re.IGNORECASE | re.UNICODE
    )
    match = pattern.search(text)
    if not match:
        return None

    full_prefix = match.group("type").strip()
    parts = re.split(r'\s{2,}', full_prefix.strip())
    if len(parts) < 2:
        doc_type = full_prefix
        doc_number = ""
    else:
        doc_type = " ".join(parts[:-1]).strip()
        doc_number = parts[-1].strip()

    completion_date = match.group("date").strip()
    completion_date = completion_date.replace(".", "/")
    institution = match.group("institution")
    country, country_en = find_country(institution)

    return {
        "document_type": doc_type,
        "document_type_en": translate_to_en(doc_type) or "",
        "document_number": doc_number,
        "completion_date": completion_date,
        "institution_name": institution,
        "institution_name_en": translate_to_en(institution) or "",
        "country": country,
        "country_en": country_en
    }


def parse_reference_cell_ua(text):
    """Розбирає комірку Excel з даними довідки про визнання документа (номер; установа; країна; дата) на окремі поля."""
    if not text or not str(text).strip():
        return {
            "reference_number": "",
            "reference_institution": "",
            "reference_institution_en": "",
            "reference_country": "",
            "reference_country_en": "",
            "reference_issue_date": "",
        }
    parts = [p.strip() for p in str(text).split(";")]
    parts += [""] * (4 - len(parts))
    reference_number, reference_institution, reference_country, reference_issue_date = parts[:4]
    reference_issue_date = reference_issue_date.replace(".", "/")
    return {
        "reference_number": reference_number,
        "reference_institution": reference_institution,
        "reference_institution_en": translate_to_en(reference_institution) or "",
        "reference_country": reference_country,
        "reference_country_en": translate_to_en(reference_country) or "",
        "reference_issue_date": reference_issue_date,
    }


def parse_recognition_cell_ua(text):
    """Розбирає комірку Excel з даними про визнання (номер сертифіката; орган; дата) на окремі поля."""
    if not text or not str(text).strip():
        return {
            "recognition_certificate_number": "",
            "recognition_issuer": "",
            "recognition_issuer_en": "",
            "recognition_date": "",
        }
    parts = [p.strip() for p in str(text).split(";")]
    parts += [""] * (3 - len(parts))
    recognition_certificate_number, recognition_issuer, recognition_date = parts[:3]
    recognition_date = recognition_date.replace(".", "/")
    return {
        "recognition_certificate_number": recognition_certificate_number,
        "recognition_issuer": recognition_issuer,
        "recognition_issuer_en": translate_to_en(recognition_issuer) or "",
        "recognition_date": recognition_date,
    }


def import_documents_preview(file_path, db):
    """Читає Excel-файл з документами про освіту, для кожного рядка знаходить студента (fuzzy_find_student), парсить документ (parse_document) та формує прев'ю-список змін (нові/оновлення) для підтвердження користувачем перед комітом у import_docs_commit."""
    wb = load_workbook(file_path)
    sheet = wb.active
    cursor = db.cursor()
    preview_rows = []
    for row_index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        fio = row[0]
        document_text = row[1]
        reference_ua_text = row[2] if len(row) > 2 else None
        recognition_ua_text = row[3] if len(row) > 3 else None

        if not fio or not document_text:
            preview_rows.append({"row_index": row_index, "error": "Пусті дані"})
            continue
        student_id, matched_name, score = fuzzy_find_student(cursor, fio)
        if not student_id:
            preview_rows.append({"row_index": row_index, "error": f"Студент не знайдений: {fio}"})
            continue
        data = parse_document(document_text)
        if not data:
            preview_rows.append({"row_index": row_index, "error": f"Не вдалося розпізнати документ: {document_text}"})
            continue

        data.update(parse_reference_cell_ua(reference_ua_text))
        data.update(parse_recognition_cell_ua(recognition_ua_text))


        cursor.execute("""
            SELECT id, document_type, document_type_en, document_number, completion_date,
                   institution_name, institution_name_en, country, country_en
            FROM education_documents WHERE student_id=?
        """, (student_id,))
        existing_doc = cursor.fetchone()

        row_data = {"row_index": row_index, "student_id": student_id, "matched_name": matched_name, "score": score, **data}

        if existing_doc:
            row_data.update({
                "status": "Оновлення", "status_class": "text-warning",
                "existing_info": f"(існує: {existing_doc['document_number'] or 'без номера'})",
                "has_document": True,
                "existing_doc_id": existing_doc["id"],
                "old_document_type": existing_doc["document_type"] or "",
                "old_document_type_en": existing_doc["document_type_en"] or "",
                "old_document_number": existing_doc["document_number"] or "",
                "old_completion_date": existing_doc["completion_date"] or "",
                "old_institution_name": existing_doc["institution_name"] or "",
                "old_institution_name_en": existing_doc["institution_name_en"] or "",
                "old_country": existing_doc["country"] or "",
                "old_country_en": existing_doc["country_en"] or "",
            })

            cursor.execute("""
                SELECT reference_number, reference_institution, reference_institution_en,
                       reference_country, reference_country_en, reference_issue_date,
                       recognition_certificate_number, recognition_issuer, recognition_issuer_en, recognition_date
                FROM foreign_education_docs WHERE education_doc_id=?
            """, (existing_doc["id"],))
            existing_foreign = cursor.fetchone()
            if existing_foreign:
                row_data.update({
                    "old_reference_number": existing_foreign["reference_number"] or "",
                    "old_reference_institution": existing_foreign["reference_institution"] or "",
                    "old_reference_institution_en": existing_foreign["reference_institution_en"] or "",
                    "old_reference_country": existing_foreign["reference_country"] or "",
                    "old_reference_country_en": existing_foreign["reference_country_en"] or "",
                    "old_reference_issue_date": existing_foreign["reference_issue_date"] or "",
                    "old_recognition_certificate_number": existing_foreign["recognition_certificate_number"] or "",
                    "old_recognition_issuer": existing_foreign["recognition_issuer"] or "",
                    "old_recognition_issuer_en": existing_foreign["recognition_issuer_en"] or "",
                    "old_recognition_date": existing_foreign["recognition_date"] or "",
                })
        else:
            row_data.update({"status": "Новий", "status_class": "text-success", "existing_info": "", "has_document": False})

        preview_rows.append(row_data)
    return preview_rows


@admin_bp.route('/admin/import_education_docs_preview', methods=['GET', 'POST'])
@permission_required('import_education_docs')
def import_docs_preview():
    """Приймає завантажений Excel-файл, будує прев'ю через import_documents_preview і показує сторінку підтвердження перед фактичним імпортом."""
    if request.method == "POST":
        file = request.files.get("file")
        if not file:
            flash("Файл не обрано", "danger")
            return redirect(url_for('admin.manage_education_documents'))

        filename = secure_filename(file.filename)
        path = os.path.join("uploads", filename)
        os.makedirs("uploads", exist_ok=True)
        file.save(path)

        db = get_db()
        preview = import_documents_preview(path, db)
        preview_id = save_preview_to_file(preview)

        return render_template("import_education_preview.html", preview=preview, preview_id=preview_id)

    return render_template("import_education_upload.html")


@admin_bp.route('/admin/import_docs_commit', methods=['POST'])
@permission_required('import_education_docs')
def import_docs_commit():
    """Записує в БД документи про освіту (та дані про визнання), підтверджені користувачем на сторінці прев'ю (тільки позначені чекбоксом рядки)."""
    preview_id = request.form.get("preview_id")
    preview = load_preview_from_file(preview_id)
    db = get_db()
    added = 0
    updated = 0

    for row in preview:
        row_index = row["row_index"]
        if f"add_{row_index}" not in request.form:
            continue

        if not row.get("error"):
            row["document_type"]       = request.form.get(f"document_type_{row_index}",       row["document_type"])
            row["document_type_en"]    = request.form.get(f"document_type_en_{row_index}",    row.get("document_type_en", "")) or ""
            row["document_number"]     = request.form.get(f"document_number_{row_index}",     row["document_number"])
            row["completion_date"]     = request.form.get(f"completion_date_{row_index}",     row["completion_date"])
            row["institution_name"]    = request.form.get(f"institution_name_{row_index}",    row["institution_name"])
            row["institution_name_en"] = request.form.get(f"institution_name_en_{row_index}", row.get("institution_name_en", "")) or ""
            row["country"]             = request.form.get(f"country_{row_index}",             row["country"])
            row["country_en"]          = request.form.get(f"country_en_{row_index}",          row["country_en"])

            reference_number = request.form.get(f"reference_number_{row_index}", row.get("reference_number", "")) or None
            reference_institution = request.form.get(f"reference_institution_{row_index}", row.get("reference_institution", "")) or None
            reference_institution_en = request.form.get(f"reference_institution_en_{row_index}", row.get("reference_institution_en", "")) or None
            reference_country = request.form.get(f"reference_country_{row_index}", row.get("reference_country", "")) or None
            reference_country_en = request.form.get(f"reference_country_en_{row_index}", row.get("reference_country_en", "")) or None
            reference_issue_date = request.form.get(f"reference_issue_date_{row_index}", row.get("reference_issue_date", "")) or None
            recognition_certificate_number = request.form.get(f"recognition_certificate_number_{row_index}", row.get("recognition_certificate_number", "")) or None
            recognition_issuer = request.form.get(f"recognition_issuer_{row_index}", row.get("recognition_issuer", "")) or None
            recognition_issuer_en = request.form.get(f"recognition_issuer_en_{row_index}", row.get("recognition_issuer_en", "")) or None
            recognition_date = request.form.get(f"recognition_date_{row_index}", row.get("recognition_date", "")) or None

            country_lower = (row["country"] or "").lower()

            is_foreign = not (country_lower in ('україна', 'ukraine'))

            cursor = db.cursor()
            cursor.execute("SELECT id FROM education_documents WHERE student_id=?", (row["student_id"],))
            existing = cursor.fetchone()

            doc_type_en  = row.get("document_type_en", "") or ""
            inst_name_en = row.get("institution_name_en", "") or ""
            country_en   = row.get("country_en", "") or ""

            if existing:
                education_doc_id = existing["id"]
                cursor.execute("""
                    UPDATE education_documents SET
                        document_type=?, document_type_en=?, document_number=?, completion_date=?,
                        institution_name=?, institution_name_en=?, country=?, country_en=?
                    WHERE id=?
                """, (row["document_type"], doc_type_en, row["document_number"], row["completion_date"],
                      row["institution_name"], inst_name_en, row["country"], row["country_en"], education_doc_id))
                updated += 1
            else:
                cursor.execute("""
                    INSERT INTO education_documents (
                        student_id, document_type, document_type_en, document_number,
                        completion_date, institution_name, institution_name_en, country, country_en
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (row["student_id"], row["document_type"], doc_type_en, row["document_number"],
                      row["completion_date"], row["institution_name"], inst_name_en, row["country"], country_en))
                education_doc_id = cursor.lastrowid
                added += 1

            # --- Foreign / recognition data ---
            cursor.execute(
                "SELECT id FROM foreign_education_docs WHERE education_doc_id = ?",
                (education_doc_id,)
            )
            foreign_exists = cursor.fetchone()

            # Если заполнена хотя бы одна колонка аккредитации/признания
            has_foreign_data = any([
                reference_number,
                reference_institution,
                reference_country,
                reference_issue_date,
                recognition_certificate_number,
                recognition_issuer,
                recognition_date
            ])

            if foreign_exists:
                if has_foreign_data:
                    cursor.execute("""
                        UPDATE foreign_education_docs SET
                            reference_number=?, reference_institution=?, reference_institution_en=?,
                            reference_country=?, reference_country_en=?, reference_issue_date=?,
                            recognition_certificate_number=?, recognition_issuer=?,
                            recognition_issuer_en=?, recognition_date=?
                        WHERE education_doc_id=?
                    """, (
                        reference_number,
                        reference_institution,
                        reference_institution_en,
                        reference_country,
                        reference_country_en,
                        reference_issue_date,
                        recognition_certificate_number,
                        recognition_issuer,
                        recognition_issuer_en,
                        recognition_date,
                        education_doc_id
                    ))
                else:
                    # если все поля пустые - удаляем старую запись
                    cursor.execute(
                        "DELETE FROM foreign_education_docs WHERE education_doc_id=?",
                        (education_doc_id,)
                    )

            else:
                if has_foreign_data:
                    cursor.execute("""
                        INSERT INTO foreign_education_docs (
                            education_doc_id,
                            reference_number,
                            reference_institution,
                            reference_institution_en,
                            reference_country,
                            reference_country_en,
                            reference_issue_date,
                            recognition_certificate_number,
                            recognition_issuer,
                            recognition_issuer_en,
                            recognition_date
                        )
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """, (
                        education_doc_id,
                        reference_number,
                        reference_institution,
                        reference_institution_en,
                        reference_country,
                        reference_country_en,
                        reference_issue_date,
                        recognition_certificate_number,
                        recognition_issuer,
                        recognition_issuer_en,
                        recognition_date
                    ))

    db.commit()

    log_action(
        current_username(),
        f"імпорт документів про освіту: додано {added}, оновлено {updated}"
    )

    flash(f"Додано записів: {added}, Оновлено записів: {updated}", "success")
    return redirect(url_for('admin.manage_education_documents'))
